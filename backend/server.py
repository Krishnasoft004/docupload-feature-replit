from fastapi import FastAPI, APIRouter, File, UploadFile, HTTPException, Form
from fastapi.responses import JSONResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field
from typing import List, Optional, Dict, Any
import uuid
from datetime import datetime, timezone
import asyncio
import json
import io
import tempfile
import numpy as np
from sklearn.metrics.pairwise import cosine_similarity
from sentence_transformers import SentenceTransformer

# Document processing imports
import pypdf
import docx
import openpyxl

# LLM integration
from emergentintegrations.llm.chat import LlmChat, UserMessage

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# Initialize sentence transformer for embeddings
embedding_model = SentenceTransformer('all-MiniLM-L6-v2')

# Create the main app without a prefix
app = FastAPI(title="FinQuery AI", description="AI-powered financial document assistant")

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")

# In-memory storage for document chunks and embeddings (MVP approach)
document_store = {}
embeddings_store = {}

# Define Models
class ChatMessage(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    session_id: str
    message: str
    response: str
    sources: Optional[List[Dict[str, Any]]] = []
    timestamp: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class ChatRequest(BaseModel):
    session_id: str
    message: str

class DocumentInfo(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    filename: str
    file_size: int
    upload_date: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    content_preview: str
    chunk_count: int

class ChatResponse(BaseModel):
    response: str
    sources: List[Dict[str, Any]] = []
    session_id: str

# Helper functions
def extract_text_from_pdf(file_content: bytes) -> str:
    """Extract text from PDF file."""
    try:
        pdf_file = io.BytesIO(file_content)
        reader = pypdf.PdfReader(pdf_file)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error processing PDF: {str(e)}")

def extract_text_from_docx(file_content: bytes) -> str:
    """Extract text from DOCX file."""
    try:
        doc_file = io.BytesIO(file_content)
        doc = docx.Document(doc_file)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error processing DOCX: {str(e)}")

def extract_text_from_xlsx(file_content: bytes) -> str:
    """Extract text from XLSX file."""
    try:
        excel_file = io.BytesIO(file_content)
        workbook = openpyxl.load_workbook(excel_file)
        text = ""
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            text += f"Sheet: {sheet_name}\n"
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                if row_text.strip():
                    text += row_text + "\n"
            text += "\n"
        return text
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error processing XLSX: {str(e)}")

def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
    """Split text into chunks with overlap."""
    chunks = []
    start = 0
    text_length = len(text)
    
    while start < text_length:
        end = start + chunk_size
        chunk = text[start:end]
        
        # Try to break at sentence boundary if possible
        if end < text_length:
            last_sentence = chunk.rfind('.')
            if last_sentence > chunk_size // 2:  # Only if we found a reasonable break point
                chunk = chunk[:last_sentence + 1]
                end = start + last_sentence + 1
        
        chunks.append(chunk.strip())
        start = end - overlap
        
        if start >= text_length:
            break
    
    return [chunk for chunk in chunks if chunk.strip()]

def find_relevant_chunks(query: str, doc_id: str, top_k: int = 3) -> List[Dict[str, Any]]:
    """Find most relevant chunks for a query using cosine similarity."""
    if doc_id not in embeddings_store:
        return []
    
    # Get query embedding
    query_embedding = embedding_model.encode([query])
    
    # Get document embeddings
    doc_embeddings = embeddings_store[doc_id]['embeddings']
    chunks = embeddings_store[doc_id]['chunks']
    
    # Calculate similarities
    similarities = cosine_similarity(query_embedding, doc_embeddings)[0]
    
    # Get top k most similar chunks
    top_indices = np.argsort(similarities)[-top_k:][::-1]
    
    relevant_chunks = []
    for idx in top_indices:
        if similarities[idx] > 0.1:  # Threshold for relevance
            relevant_chunks.append({
                'text': chunks[idx],
                'similarity': float(similarities[idx]),
                'chunk_index': int(idx)
            })
    
    return relevant_chunks

async def generate_rag_response(query: str, session_id: str, relevant_chunks: List[Dict[str, Any]]) -> str:
    """Generate response using RAG with relevant document chunks."""
    # Prepare context from relevant chunks
    context = ""
    if relevant_chunks:
        context = "Based on the following financial document excerpts:\n\n"
        for i, chunk in enumerate(relevant_chunks, 1):
            context += f"Source {i}:\n{chunk['text']}\n\n"
        context += f"Question: {query}\n\nPlease provide a detailed answer based only on the information provided in the sources above. If the answer cannot be found in the sources, please say so clearly."
    else:
        context = f"I don't have any relevant financial documents uploaded to answer the question: {query}\n\nPlease upload relevant financial documents first to get accurate answers."
    
    # Initialize LLM chat
    chat = LlmChat(
        api_key=os.environ.get('EMERGENT_LLM_KEY'),
        session_id=session_id,
        system_message="You are FinQuery AI, a specialized financial document assistant. You help users analyze financial documents by providing accurate answers based solely on the uploaded document content. Always cite your sources and be precise about financial data."
    ).with_model("gemini", "gemini-2.0-flash")
    
    # Create user message
    user_message = UserMessage(text=context)
    
    # Get response from LLM
    response = await chat.send_message(user_message)
    return response

# API Routes
@api_router.post("/upload-document", response_model=DocumentInfo)
async def upload_document(file: UploadFile = File(...)):
    """Upload and process a financial document."""
    
    # Check file size (20MB limit)
    MAX_FILE_SIZE = 20 * 1024 * 1024  # 20MB
    
    # Read file content
    file_content = await file.read()
    
    if len(file_content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="File size exceeds 20MB limit")
    
    # Check file type and extract text
    file_extension = Path(file.filename).suffix.lower()
    
    if file_extension == '.pdf':
        text = extract_text_from_pdf(file_content)
    elif file_extension == '.docx':
        text = extract_text_from_docx(file_content)
    elif file_extension == '.xlsx':
        text = extract_text_from_xlsx(file_content)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file type. Please upload PDF, DOCX, or XLSX files.")
    
    if not text.strip():
        raise HTTPException(status_code=400, detail="No text content found in the document")
    
    # Generate chunks
    chunks = chunk_text(text)
    
    if not chunks:
        raise HTTPException(status_code=400, detail="Could not extract meaningful content from document")
    
    # Generate embeddings
    embeddings = embedding_model.encode(chunks)
    
    # Generate document ID
    doc_id = str(uuid.uuid4())
    
    # Store document chunks and embeddings in memory
    document_store[doc_id] = {
        'filename': file.filename,
        'chunks': chunks,
        'full_text': text,
        'upload_date': datetime.now(timezone.utc)
    }
    
    embeddings_store[doc_id] = {
        'embeddings': embeddings,
        'chunks': chunks
    }
    
    # Store document info in database
    doc_info = DocumentInfo(
        id=doc_id,
        filename=file.filename,
        file_size=len(file_content),
        content_preview=text[:500] + "..." if len(text) > 500 else text,
        chunk_count=len(chunks)
    )
    
    await db.documents.insert_one(doc_info.dict())
    
    return doc_info

@api_router.post("/chat", response_model=ChatResponse)
async def chat_with_documents(chat_request: ChatRequest):
    """Ask questions about uploaded documents."""
    
    # Find relevant chunks from all uploaded documents
    all_relevant_chunks = []
    
    for doc_id in document_store.keys():
        relevant_chunks = find_relevant_chunks(chat_request.message, doc_id, top_k=2)
        for chunk in relevant_chunks:
            chunk['document_id'] = doc_id
            chunk['filename'] = document_store[doc_id]['filename']
        all_relevant_chunks.extend(relevant_chunks)
    
    # Sort by similarity and take top chunks
    all_relevant_chunks.sort(key=lambda x: x['similarity'], reverse=True)
    top_chunks = all_relevant_chunks[:5]  # Top 5 most relevant chunks
    
    # Generate response using RAG
    response = await generate_rag_response(chat_request.message, chat_request.session_id, top_chunks)
    
    # Prepare sources for response
    sources = []
    for chunk in top_chunks:
        sources.append({
            'filename': chunk['filename'],
            'text': chunk['text'][:200] + "..." if len(chunk['text']) > 200 else chunk['text'],
            'similarity': chunk['similarity']
        })
    
    # Store chat message in database
    chat_message = ChatMessage(
        session_id=chat_request.session_id,
        message=chat_request.message,
        response=response,
        sources=sources
    )
    
    await db.chat_messages.insert_one(chat_message.dict())
    
    return ChatResponse(
        response=response,
        sources=sources,
        session_id=chat_request.session_id
    )

@api_router.get("/documents", response_model=List[DocumentInfo])
async def get_documents():
    """Get list of uploaded documents."""
    documents = await db.documents.find().to_list(100)
    return [DocumentInfo(**doc) for doc in documents]

@api_router.get("/chat-history/{session_id}")
async def get_chat_history(session_id: str):
    """Get chat history for a session."""
    chat_history = await db.chat_messages.find({"session_id": session_id}).to_list(100)
    return [ChatMessage(**msg) for msg in chat_history]

@api_router.get("/")
async def root():
    return {"message": "FinQuery AI Backend Running", "status": "operational"}

@api_router.get("/health")
async def health_check():
    """Health check endpoint."""
    return {
        "status": "healthy",
        "documents_loaded": len(document_store),
        "embedding_model": "all-MiniLM-L6-v2",
        "llm_model": "gemini-2.0-flash"
    }

# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()